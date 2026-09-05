from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Brand_Visibility_Intelligence_Step_by_Step_Guide.docx"
BLUE = "2E5BFF"
NAVY = "172033"
PALE = "E8EEF5"
LIGHT = "F4F6F9"
MUTED = "667085"
WHITE = "FFFFFF"

CODE_FILES = [
    "pyproject.toml", "requirements.txt", ".env.example", ".gitignore", "Makefile", "Dockerfile",
    ".dockerignore", ".streamlit/config.toml", "config/logging.ini",
    "src/brand_visibility/__init__.py", "src/brand_visibility/config.py",
    "src/brand_visibility/pipeline.py", "src/brand_visibility/data/extract.py",
    "src/brand_visibility/data/sample_data.py", "src/brand_visibility/data/transform.py",
    "src/brand_visibility/data/database.py", "src/brand_visibility/analytics/eda.py",
    "src/brand_visibility/dashboard/app.py", "scripts/export_eda.py", "tests/test_transform.py",
]


def font(run, name="Aptos", size=10.5, color=NAVY, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_toc(doc):
    entries = [
        "1. Project objective and deliverables", "2. Architecture and folder structure",
        "3. Environment setup", "4. Data extraction", "5. CSV integration and schema alignment",
        "6. Data cleaning", "7. Feature engineering", "8. SQLite serving layer",
        "9. Exploratory data analysis", "10. Dashboard operation", "11. Testing and quality checks",
        "12. Docker and production operation", "13. Troubleshooting",
        "Appendix A. Complete source code", "Appendix B. Command reference",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(3)
        font(p.add_run(entry), size=10, color=NAVY, bold=entry.startswith("Appendix"))


def add_callout(doc, label, text, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.32)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}: ")
    font(r, bold=True, color=BLUE)
    font(p.add_run(text))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, code, size=7.4):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = False
    p.paragraph_format.widow_control = False
    # Tabs are normalized to four spaces so code is stable in Word and LibreOffice.
    for index, line in enumerate(code.expandtabs(4).splitlines() or [""]):
        r = p.add_run(line)
        font(r, "Consolas", size, "202938")
        if index < len(code.splitlines()) - 1:
            r.add_break()
    return p


def add_command(doc, command):
    p = doc.add_paragraph(style="Command")
    font(p.add_run(command), "Consolas", 8.2, WHITE)


def add_steps(doc, steps):
    for title, detail in steps:
        p = doc.add_paragraph(style="Step")
        r = p.add_run(title)
        font(r, bold=True, color=NAVY)
        p2 = doc.add_paragraph(detail)
        p2.paragraph_format.left_indent = Inches(0.38)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after in (("Title", 28, 0, 12), ("Heading 1", 17, 18, 9), ("Heading 2", 13.5, 14, 7), ("Heading 3", 11.5, 10, 5)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor.from_string(BLUE if name != "Heading 3" else NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Subtitle"].font.name = "Aptos"
    styles["Subtitle"].font.size = Pt(14)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(7.4)
    code.paragraph_format.left_indent = Inches(0.08)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0
    ppr = code.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F3F6")
    ppr.append(shd)

    command = styles.add_style("Command", WD_STYLE_TYPE.PARAGRAPH)
    command.font.name = "Consolas"
    command.font.size = Pt(8.2)
    command.font.color.rgb = RGBColor.from_string(WHITE)
    command.paragraph_format.left_indent = Inches(0.08)
    command.paragraph_format.right_indent = Inches(0.08)
    command.paragraph_format.space_before = Pt(4)
    command.paragraph_format.space_after = Pt(8)
    command.paragraph_format.line_spacing = 1.05
    ppr = command.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), NAVY)
    ppr.append(shd)

    step = styles.add_style("Step", WD_STYLE_TYPE.PARAGRAPH)
    step.font.name = "Aptos"
    step.font.size = Pt(11)
    step.paragraph_format.left_indent = Inches(0.38)
    step.paragraph_format.first_line_indent = Inches(-0.38)
    step.paragraph_format.space_before = Pt(8)
    step.paragraph_format.space_after = Pt(2)


def add_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = "BRAND VISIBILITY INTELLIGENCE  |  IMPLEMENTATION GUIDE"
    font(header.runs[0], size=8, color=MUTED, bold=True)
    header.paragraph_format.space_after = Pt(2)
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = t.rows[0].cells[i]
        cell.width = Inches(width)
        shade(cell, PALE)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        font(p.add_run(header), size=9, bold=True, color=NAVY)
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, (value, width) in enumerate(zip(row, widths)):
            cells[i].width = Inches(width)
            cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            font(cells[i].paragraphs[0].add_run(str(value)), size=8.8)
    return t


def build():
    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.49)
    add_header_footer(sec)

    # Editorial-cover opening for a long technical manual.
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph("END-TO-END PROJECT MANUAL")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(kicker.add_run(), color=BLUE)
    title = doc.add_paragraph("Brand Visibility\nIntelligence Dashboard", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Step-by-step implementation, complete source code, validation and operations", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    meta = doc.add_paragraph(f"Version 1.0  |  {date.today().strftime('%d %B %Y')}  |  Python · SQLite · Streamlit · Plotly")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(meta.runs[0], size=9.5, color=MUTED, bold=True)
    doc.add_page_break()

    doc.add_heading("How to use this guide", level=1)
    doc.add_paragraph("This manual is both a learning path and a technical reference. Follow Chapters 1-12 in order for a first implementation. Use the appendices when copying or reviewing the complete project source.")
    add_callout(doc, "Recommended workflow", "Create the environment, run the demo pipeline, verify the dashboard, then replace demo input with the supplied dirty CSV and finally enable SerpAPI.")
    doc.add_heading("Table of contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    doc.add_heading("1. Project objective and deliverables", level=1)
    doc.add_paragraph("The system measures how products and brands perform in shopping search results. It combines external API results with an existing CSV, cleans and enriches the records, stores them in SQLite, answers 30 analytical questions and serves an interactive dashboard.")
    table(doc, ["Deliverable", "Implementation"], [
        ("Source code", "Installable Python package under src/brand_visibility"),
        ("Clean dataset", "data/processed/cleaned_products.csv"),
        ("SQL serving layer", "data/database/brand_visibility.db"),
        ("Dashboard", "Six-tab Streamlit application with dynamic SQL filters"),
        ("EDA report", "reports/eda_answers.md with all 30 questions"),
        ("Quality controls", "Tests, lint configuration, Docker health check and cleaning report"),
    ], [1.65, 4.67])

    doc.add_heading("2. Architecture and folder structure", level=1)
    add_code(doc, """Brand/
├── config/                  # Logging and runtime configuration
├── data/{raw,interim,processed,database}/
├── docs/                    # Architecture and this manual
├── reports/                 # EDA, cleaning diagnostics and figures
├── scripts/                 # Operational scripts
├── src/brand_visibility/    # Installable application package
│   ├── analytics/           # EDA and business insight logic
│   ├── dashboard/           # Streamlit presentation layer
│   ├── data/                # Extract, transform, sample and storage
│   ├── config.py            # Central settings
│   └── pipeline.py          # ETL orchestration
└── tests/                   # Automated tests""")
    doc.add_paragraph("The pipeline and dashboard are intentionally separate. API or ingestion failures cannot take down the dashboard because the UI reads only the serving database.")

    doc.add_heading("3. Environment setup", level=1)
    add_steps(doc, [
        ("Step 1 - Open the project", "Start a terminal and change into the repository root."),
        ("Step 2 - Create an isolated environment", "A virtual environment prevents dependency conflicts with other projects."),
        ("Step 3 - Install the package", "Editable installation exposes brand_visibility while preserving live source edits."),
    ])
    add_command(doc, 'cd "/path/to/Brand"\npython3 -m venv .venv\nsource .venv/bin/activate\nmake install')
    doc.add_heading("Configuration", level=2)
    doc.add_paragraph("Copy `.env.example` to `.env` when using SerpAPI. Never commit the real key.")
    add_command(doc, "cp .env.example .env")
    add_code(doc, (ROOT / ".env.example").read_text())

    doc.add_heading("4. Data extraction", level=1)
    doc.add_paragraph("The extraction layer sends one Google Shopping request per keyword, validates the HTTP response and maps the external response into the internal product schema. A missing key fails early with a clear message.")
    add_callout(doc, "Security", "The API key is read from SERPAPI_KEY. It is never hard-coded in the source or written into exported datasets.")
    add_command(doc, "python -m brand_visibility.pipeline --api --keywords laptop phone headphones smartwatch")
    doc.add_heading("Important extraction logic", level=2)
    extract_text = (ROOT / "src/brand_visibility/data/extract.py").read_text()
    add_code(doc, extract_text)

    doc.add_heading("5. CSV integration and schema alignment", level=1)
    doc.add_paragraph("CSV files from real systems rarely use identical headings. `align_columns` normalizes case and punctuation, maps common aliases, creates absent fields and returns a canonical column order before concatenation.")
    add_steps(doc, [
        ("Step 1 - Load", "Pandas reads the supplied brand_dirty_dataset.csv."),
        ("Step 2 - Normalize", "Column names become lower snake_case and aliases are mapped."),
        ("Step 3 - Complete", "Missing canonical fields are created with null values."),
        ("Step 4 - Combine", "API and CSV rows are appended into one dataframe."),
    ])
    add_command(doc, "python -m brand_visibility.pipeline --csv /path/to/brand_dirty_dataset.csv")

    doc.add_heading("6. Data cleaning", level=1)
    table(doc, ["Issue", "Treatment"], [
        ("Price/reviews stored as text", "Remove currency/noise characters; safely coerce to numeric"),
        ("Invalid price or rank", "Treat non-positive values as missing"),
        ("Invalid rating", "Accept only values from 0 to 5"),
        ("Missing numeric values", "Keyword median, then dataset median"),
        ("Missing delivery", "Label Unknown"),
        ("Noisy product titles", "Remove unwanted symbols and collapse whitespace"),
        ("Platform casing", "Normalize to title case"),
        ("Extreme price", "Cap at the 99th percentile and record the threshold"),
        ("Duplicates", "Remove repeated keyword + title + platform + price records"),
    ], [2.0, 4.32])
    add_callout(doc, "Business rule", "Products without ratings or reviews are retained because their price, platform and search visibility still contribute useful market evidence.")

    doc.add_heading("7. Feature engineering", level=1)
    table(doc, ["Feature", "Definition"], [
        ("brand", "Known brand found in title; otherwise a conservative first-token fallback"),
        ("visibility_score", "100 / position, so rank 1 receives 100"),
        ("discount_pct", "(raw_price - price) / raw_price × 100, clipped to 0-100"),
        ("price_range", "Under $50, $50-$149, $150-$499, $500-$999, or $1,000+"),
        ("rating_range", "Under 3, 3-3.9, 4-4.4, or 4.5+"),
        ("top_10", "1 when position ≤ 10; otherwise 0"),
    ], [1.65, 4.67])

    doc.add_heading("8. SQLite serving layer", level=1)
    doc.add_paragraph("The cleaned dataframe replaces the products table in one controlled operation. A composite index accelerates dashboard filters. All user filter values are passed as SQL parameters, which avoids SQL injection and quoting bugs.")
    add_command(doc, "make pipeline")
    doc.add_paragraph("Generated outputs:", style=None)
    for item in ["data/processed/cleaned_products.csv", "data/database/brand_visibility.db", "reports/cleaning_report.json"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("9. Exploratory data analysis", level=1)
    doc.add_paragraph("The analysis module returns one named result for every question in the brief. Results include grouped summaries, distributions, top products, cross-tabs and correlations.")
    table(doc, ["Group", "Questions covered"], [
        ("General market", "Products by keyword; price/rating/review distributions"),
        ("Brand", "Frequency, visibility, ranking, top-10 presence and ratings"),
        ("Pricing", "Brand distributions, ranges, price/rank relationship and maxima"),
        ("Discounts", "Discount adoption, ranking, brand/platform averages and rating relationship"),
        ("Platforms", "Volume, rating, price, ranking and brand mix"),
        ("Visibility", "Ranking distribution, visibility leaders and factor relationships"),
    ], [1.55, 4.77])
    add_command(doc, "make eda")
    add_callout(doc, "Interpretation", "Correlation indicates association, not causation. Search rank may also be affected by relevance, inventory, advertising and personalization factors that are not present in the dataset.")

    doc.add_heading("10. Dashboard operation", level=1)
    add_command(doc, "make dashboard")
    doc.add_paragraph("Open http://localhost:8501. Sidebar selections are translated into parameterized SQLite queries and refresh all tabs.")
    table(doc, ["Tab", "Purpose"], [
        ("Overview", "Core KPIs, price distribution, keyword volume, platform share and insights"),
        ("Brand Insights", "Brand volume, ratings, visibility and top-10 presence"),
        ("Pricing Analysis", "Price ranges, rank/rating relationships and discounts"),
        ("Platform Analysis", "Platform volume, price, rating, position and brand distribution"),
        ("Visibility & Ranking", "Ranking distribution, correlations and visibility leaders"),
        ("Product Explorer", "Search, sortable top-10 table, charts and CSV download"),
    ], [1.65, 4.67])

    doc.add_heading("11. Testing and quality checks", level=1)
    add_command(doc, "make test\nmake lint")
    doc.add_paragraph("The current transformation test verifies numeric cleaning, duplicate removal, brand extraction, visibility calculation and discount calculation. For production, extend tests for empty datasets, API errors, schema drift and boundary values.")
    checklist = [
        "Pipeline exits successfully and records the cleaning diagnostics.",
        "Clean numeric fields contain no missing values.",
        "SQLite row count matches the processed CSV.",
        "All 30 EDA keys are present.",
        "Dashboard filters return valid subsets and empty selections are handled.",
        "Secrets are absent from source control.",
    ]
    for item in checklist:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("12. Docker and production operation", level=1)
    add_command(doc, "docker build -t brand-visibility .\ndocker run --rm -p 8501:8501 brand-visibility")
    doc.add_paragraph("The image installs the project package, exposes port 8501 and uses Streamlit's health endpoint. In a larger deployment, mount durable data storage, inject secrets through the platform's secret manager and run the pipeline as a scheduled job separate from the dashboard container.")
    add_callout(doc, "Production recommendation", "Use managed PostgreSQL when multiple dashboard replicas or concurrent writers are required. SQLite is appropriate for a single-instance learning project and small read-heavy deployments.")

    doc.add_heading("13. Troubleshooting", level=1)
    table(doc, ["Symptom", "Resolution"], [
        ("No module named brand_visibility", "Activate the environment and run `make install` from the root"),
        ("SERPAPI_KEY is required", "Copy .env.example to .env and add a valid key, or omit --api"),
        ("Dashboard has no records", "Run `make pipeline`; check sidebar filters and database path"),
        ("Unexpected CSV columns", "Add a source alias in ALIASES or update the canonical schema"),
        ("Charts are empty after filtering", "Widen position/rating/brand selections"),
        ("Port 8501 already in use", "Run Streamlit with `--server.port 8502`"),
    ], [2.15, 4.17])

    # Named Code Appendix override: landscape pages, 0.65-inch margins, 7.1 pt Consolas.
    section = doc.add_section()
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    section.header_distance = section.footer_distance = Inches(0.35)
    add_header_footer(section)
    doc.add_heading("Appendix A. Complete source code", level=1)
    doc.add_paragraph("The following listings reproduce every authored code and configuration file required to build, run, test and package the project. Generated CSV, SQLite and report outputs are intentionally excluded.")
    for index, relative in enumerate(CODE_FILES, start=1):
        path = ROOT / relative
        if not path.exists():
            continue
        if index > 1:
            doc.add_page_break()
        doc.add_heading(f"A.{index}  {relative}", level=2)
        add_code(doc, path.read_text(), size=7.1)

    doc.add_page_break()
    doc.add_heading("Appendix B. Command reference", level=1)
    table(doc, ["Task", "Command"], [
        ("Install", 'python -m pip install -e ".[dev]"'),
        ("Demo pipeline", "python -m brand_visibility.pipeline"),
        ("CSV pipeline", "python -m brand_visibility.pipeline --csv /path/file.csv"),
        ("Live API", "python -m brand_visibility.pipeline --api --keywords laptop phone"),
        ("Dashboard", "streamlit run src/brand_visibility/dashboard/app.py"),
        ("EDA export", "python scripts/export_eda.py"),
        ("Tests", "pytest -q"),
        ("Lint", "ruff check src tests scripts"),
        ("Docker", "docker build -t brand-visibility ."),
    ], [1.65, 7.55])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Brand Visibility Intelligence Dashboard - Step-by-Step Guide"
    doc.core_properties.subject = "Implementation manual and complete source code"
    doc.core_properties.author = "Brand Visibility Intelligence Project"
    doc.core_properties.keywords = "Python, Streamlit, SQLite, SerpAPI, EDA, Dashboard"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
